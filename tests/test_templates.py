from pathlib import Path
from xml.etree import ElementTree

import pytest
import sbol2

from excel2sbol.compiler import initialise, initialise_welcome
from excel2sbol.converter import converter


TEMPLATES_DIR = Path(__file__).parents[1] / "resources" / "templates"
NIST_WORKBOOKS_DIR = Path(__file__).parents[1] / "resources" / "NIST_workbooks"
NIST_WORKBOOK_OUTPUTS_DIR = Path(__file__).parent / "NIST_workbook_outputs"

TEMPLATE_SHEETS = {
    "Base.xlsm": {"welcome", "ontology_terms", "organism_terms"},
    "Resources.xlsm": {"Init", "column_definitions", "chassis", "cds"},
    "SampleDesign.xlsm": {"Init", "column_definitions", "sample design", "supplement"},
    "Strains.xlsm": {"Init", "column_definitions", "strain"},
    "Study.xlsm": {"Init", "column_definitions", "study", "assay", "sample"},
}

# Base is a source workbook and Resources currently includes a workbook-only
# "Translate to Protein" field. The remaining templates are complete converter
# inputs and must produce valid SBOL.
CONVERTIBLE_TEMPLATES = ("SampleDesign.xlsm", "Strains.xlsm", "Study.xlsm")
NIST_WORKBOOKS = tuple(sorted(path.name for path in NIST_WORKBOOKS_DIR.glob("*.xlsm")))


@pytest.mark.parametrize("template_name", TEMPLATE_SHEETS)
def test_template_contains_required_sheets(template_name):
    """Ensure every distributed template remains a valid Excel workbook."""
    import openpyxl

    workbook = openpyxl.load_workbook(TEMPLATES_DIR / template_name, read_only=True)

    assert TEMPLATE_SHEETS[template_name].issubset(workbook.sheetnames)


@pytest.mark.parametrize("template_name", sorted(set(TEMPLATE_SHEETS) - {"Base.xlsm"}))
def test_converter_can_read_template_configuration(template_name):
    """Ensure templates can be read by the installed converter package."""
    col_definitions, sheets_to_convert, compiled_sheets, version, _, init_info = initialise(
        TEMPLATES_DIR / template_name
    )

    assert version == 2
    assert sheets_to_convert
    assert set(sheets_to_convert).issubset(compiled_sheets)
    assert not col_definitions.empty
    assert initialise_welcome(init_info, TEMPLATES_DIR / template_name) is not None


@pytest.mark.parametrize("template_name", CONVERTIBLE_TEMPLATES)
def test_converter_generates_valid_sbol(template_name, tmp_path):
    """Convert each standalone template and validate the resulting SBOL2 XML."""
    output_path = tmp_path / f"{Path(template_name).stem}.xml"

    converter(file_path_in=TEMPLATES_DIR / template_name, file_path_out=output_path)

    assert output_path.is_file()
    assert output_path.stat().st_size > 0
    ElementTree.parse(output_path)

    document = sbol2.Document()
    document.read(str(output_path))
    assert document.validate() == "Valid."


@pytest.mark.parametrize("workbook_name", NIST_WORKBOOKS)
def test_converter_generates_valid_sbol_from_nist_workbook(workbook_name):
    """Convert each NIST workbook, retain its XML output, and validate it."""
    NIST_WORKBOOK_OUTPUTS_DIR.mkdir(exist_ok=True)
    output_path = NIST_WORKBOOK_OUTPUTS_DIR / f"{Path(workbook_name).stem}.xml"

    converter(file_path_in=NIST_WORKBOOKS_DIR / workbook_name, file_path_out=output_path)

    assert output_path.is_file()
    assert output_path.stat().st_size > 0
    ElementTree.parse(output_path)

    document = sbol2.Document()
    document.read(str(output_path))
    assert document.validate() == "Valid."
